# -*- coding: utf-8 -*-
"""Chiroyli yozma savolnoma "Bilim sinovi" (DOCX + PDF): muqova, bo'lim bannerlari, savol kartochkalari,
javob maydonlari va qo'shimcha kreativ topshiriqlar (vaziyat, to'g'ri/noto'g'ri, moslashtirish, hisob jadvali).
Manba: Evalution_sheet_Abdumalikov_Ulugbek.xlsx.   python3 build_bilim_sinovi.py"""
import subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY, GOLD, INK, MUTED, PALE, LINE = '14213D', 'E0A443', '1C2430', '6B7380', 'F6F3EC', 'C9CED6'
HEAD, BODY = 'Montserrat', 'Inter'
C = lambda h: RGBColor.from_string(h)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hex_); tcPr.append(s)


def borders(cell, **kw):
    """kw: top/bottom/left/right = (size_eighths, hex) or None for nil."""
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}')
        v = kw.get(side)
        if v is None: e.set(qn('w:val'), 'nil')
        else: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(v[0])); e.set(qn('w:color'), v[1])
        b.append(e)
    tcPr.append(b)


def margins(cell, t=80, b=80, l=140, r=140):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for k, v in (('top', t), ('bottom', b), ('start', l), ('end', r)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)


def run(p, text, size=10.5, bold=False, color=INK, font=BODY, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = C(color)
    r.font.name = font; r._element.rPr.rFonts.set(qn('w:eastAsia'), font); r._element.rPr.rFonts.set(qn('w:hAnsi'), font)
    return r


def para(container, space_after=0, align=None):
    p = container.add_paragraph(); p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    if align: p.alignment = align
    return p


def table(doc, rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) * 567))); tw.set(qn('w:type'), 'dxa')
    for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
    tblPr.append(tw)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(w * 567)))
    for i, w in enumerate(widths): t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t


def gap(doc, pt=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(''); r.font.size = Pt(pt / 2); p.paragraph_format.line_spacing = Pt(pt)


W = 17.0  # foydali kenglik, sm

# ---------------- savollar ----------------
# (turi, teg, matn, javob_qatorlari)  turi: q — oddiy savol, v — vaziyatli masala
SECTIONS = [
 ('A', 'Talab etiladigan xabardorlik', 'Qadriyatlar, rol va javobgarlik, xavfsizlik, intizom', [
  ('q', 'PE-1', "Korxonaning qadriyatlari, missiyasi va strategik maqsadlari kundalik ishingizda qanday namoyon bo'ladi? Ikkita amaliy misol keltiring.", 5),
  ('q', 'PE-7', "R&R (Rol va Javobgarliklar) nima va nima uchun tashkilotning har bir darajasi uchun belgilanadi? R&R ADM MS ning nechta tamoyilini qamrab olishi kerak?", 4),
  ('q', 'PE-7', "O'z lavozimingiz bo'yicha asosiy majburiyatlaringiz, huquqlaringiz va javobgarliklaringizni yozing.", 6),
  ('q', 'PE-7', "Guruh rahbari (GL) uchun belgilangan 4 ta asosiy vazifani sanab bering.", 4),
  ('q', 'PE-2', "Ish joyi xavfsizligi tizimi (WSS) nimaga xizmat qiladi? “Har bir Inson, Har bir Ob'ekt, Har Kuni” tamoyilini qanday tushunasiz?", 4),
  ('q', 'PE-2', "Risklarni baholash uchun qaysi vosita qo'llanilishi shart? ESCP tartibida muammo 24 soatda, 7 kunda va 30 kunda hal etilmasa nima qilinadi?", 5),
  ('q', 'PE-2', "Xavfsizlik aylanmalari qanday turlarga bo'linadi? Sexda rioya qilinishi shart bo'lgan tanqidiy xavfsizlik tashabbuslariga 3 ta misol keltiring.", 5),
  ('q', 'PE-3', "GCTA (Global Call to Actions) nima? SIF qanday hodisa? Ular bo'yicha tuzatuvchi chora-tadbirlar rejasi (CAP) nimalarni o'z ichiga oladi va qanday tasdiqlanadi?", 6),
  ('q', 'PE-6', "JIT (Job Instruction Training) nima va uni kim qo'llaydi? 4 kvadrantli ko'nikma jadvali nimani ko'rsatadi va xodim operatsiyani 3 oy bajarmasa nima o'zgaradi?", 5),
  ('q', 'Intizom', "Ichki tartib-qoidalar va mehnat intizomi bo'yicha xodimdan talab qilinadigan asosiy qoidalarni yozing.", 5),
  ('v', 'VAZIYAT · PE-2/3', "Smena boshida presslash uskunasining himoya to'sig'i shikastlanganini ko'rdingiz. Usta “reja yonib ketyapti, keyin tuzatamiz” demoqda. Harakatlaringizni qadam-baqadam yozing: kimga, qanday tartibda xabar berasiz va muammo qaysi muddatlarda qaysi darajaga ko'tariladi?", 7),
 ]),
 ('B', 'Ish joyi va standartlashtirilgan ish', '5S, vizual boshqaruv, SOS va JES', [
  ('q', 'ST-3', "5S bosqichlarini tartib bilan yozing va har birini bir jumlada izohlang.", 6),
  ('q', 'ST-3', "“Red tag” jarayoni va “shadow board” nima uchun kerak? Doskalarda holat qanday umumiy belgi bilan ko'rsatiladi?", 4),
  ('q', 'ST-3', "7 turdagi yo'qotishni sanab bering. WPO orqali ulardan birini bartaraf etishga misol keltiring.", 5),
  ('q', 'ST-6', "Siklik standartlashtirilgan ish hujjatida qanday ma'lumotlar bo'lishi kerak (kamida 5 ta)? Ish elementlari sikl vaqtining taxminan necha foizi hajmida bo'laklanadi?", 5),
  ('q', 'ST-8', "Operator standartlashtirilgan ishga amal qilayotganini qanday tekshirasiz? SOSdagi umumiy sikl vaqti real holatdan necha foizdan ko'p farq qilmasligi kerak?", 5),
  ('v', 'VAZIYAT · ST-8', "Bir operatorni 3 sikl kuzatdingiz: u JESdagi 4-element bilan 5-elementning o'rnini almashtirib bajaryapti va sikl vaqti SOSdagidan 8% uzun. Bu nimani anglatadi va TL sifatida qanday choralar ko'rasiz?", 6),
 ]),
 ('C', "Ishlab chiqarish ko'rsatkichlari", 'Takt Time, Cycle Time, HPV, UPH, Downtime', [
  ('q', 'ST-4', "Takt Time nima va qanday hisoblanadi? Masala: smena 8 soat, tanaffuslar jami 30 daqiqa, smenalik reja 225 avtomobil. Takt Time necha soniya?", 5),
  ('q', 'Tahlil', "Ishlab chiqarish rejasidagi avtomobillarni Takt Time bo'yicha tahlil qilishning 4 ta afzalligini yozing.", 4),
  ('q', 'Tahlil', "Cycle Time va Line Balance nima? Masala: Cycle Time = 120 soniya, bitta stansiya operatsiyasi 150 soniya. Qanday muammo yuzaga keladi va oqibatlari qanday?", 5),
  ('q', 'Tahlil', "HPV oshishiga qaysi omillar sabab bo'ladi (kamida 5 ta)? Ishlab chiqarish hajmi oshganda HPV odatda qanday o'zgaradi va nega?", 5),
  ('q', 'Downtime', "Rejalashtirilgan va rejalashtirilmagan Downtime farqi nima? Har biriga 3 tadan misol keltiring. Downtime nima uchun o'lchanadi?", 5),
  ('q', 'Tahlil', "I/Ch sexlarining absentizm ko'rsatkichi nima uchun jamlanib rahbariyatga taqdim etiladi? Kamida 4 ta sabab yozing.", 4),
  ('q', 'Tahlil', "#630, #690, #700 eventlar tahlili nimaga xizmat qiladi? “04” avtomobillar nima va ular bo'yicha qanday choralar ko'rasiz?", 5),
  ('q', 'Tahlil', "Oylik ishlab chiqarilgan avtomobillarni brend va model kesimida tahlil qilish qaysi 4 yo'nalishda foyda beradi?", 5),
  ('v', 'VAZIYAT · Absentizm', "So'nggi 4 haftada Bo'yash sexida dushanba kunlari ishga kelmaslik boshqa kunlarga nisbatan 2 baravar yuqori. Qanday ma'lumotlarni yig'asiz, qanday sabablarni tekshirasiz va rahbariyatga qanday choralar taklif qilasiz?", 6),
 ]),
 ('D', 'Biznes reja va takomillashtirish', 'SPQRC, BPD, SMART, PDCA, Error Proofing', [
  ('q', 'SPQRC', "SPQRC(E) yo'nalishlarining har bir harfi nimani bildiradi va ular kundalik boshqaruvda nima uchun muhim?", 4),
  ('q', 'CI-1', "BPD nima? Yagona BPD formati qanday elementlarni o'z ichiga oladi va korxonada odatda nechta BPD darajasi bo'ladi?", 5),
  ('q', 'CI-2', "Metrika doimiy “yashil” yoki doimiy “qizil” bo'lsa, nima qilish kerak?", 4),
  ('q', 'PDCA', "PDCA siklining bosqichlarini yozing va o'z ishingizdagi bitta muammoni shu sikl bo'yicha qanday hal qilishingizni ko'rsating.", 6),
  ('q', 'Q-10 · PE-15/16', "Error Proofing (xatolarning oldini olish) nima va uning validatsiyasi qanday o'tkaziladi? “Go and See” va Layered Audit nima?", 5),
  ('q', 'MS Office', "Excel'da reja-fakt tahlili uchun qaysi formulalardan foydalanasiz? Ikkita misol yozing.", 3),
  ('v', 'TAKLIF · Kaizen', "O'z ish joyingiz uchun bitta takomillashtirish taklifini yozing: muammo, sababi, yechim, kutilayotgan foyda (raqam bilan) va natijani qanday o'lchaysiz.", 7),
 ]),
]

TF = [  # to'g'ri / noto'g'ri
 "Takt Time mavjud ish vaqtini talab qilingan mahsulot soniga bo'lish orqali topiladi.",
 "Cycle Time Takt Timedan katta bo'lsa, liniya rejani bemalol bajaradi.",
 "HPV qancha past bo'lsa, mehnat samaradorligi shuncha yuqori.",
 "Tanaffus va profilaktik ta'mirlash rejalashtirilmagan Downtime hisoblanadi.",
 "JES har doim SOS bilan bog'langan bo'lishi kerak.",
 "5S ning oxirgi bosqichi — “Odat qilish” (intizom).",
 "SOSdagi umumiy sikl vaqti real holatdan 15% gacha farq qilishi mumkin.",
 "Muammo 30 kun ichida hal etilmasa, u SRB darajasiga ko'tariladi.",
]
MATCH_L = ['SOS', 'JES', 'STS', 'TIS', 'BPD', 'CAP']
MATCH_R = ['Tuzatuvchi chora-tadbirlar rejasi', 'Standart operatsiyalar varagʻi', 'Vazifa koʻrsatmalari varagʻi (siklik boʻlmagan ish tafsiloti)',
           'Biznes rejani boshqarish', 'Ish elementi yoʻriqnomasi (nima / qanday / nima uchun)', 'Standart vazifalar varagʻi (siklik boʻlmagan ish)']
CALC = [('Takt Time', 'Smena 450 daqiqa, reja 270 avtomobil'), ('HPV', '5 400 soat, 360 avtomobil'),
        ('UPH', '7,5 soatda 420 avtomobil'), ('HPV yaxshilanishi', '20 soatdan 17 soatga'), ('Reja bajarilishi', 'Reja 300, fakt 282')]

# ---------------- hujjat ----------------
d = Document()
sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.6)
st = d.styles['Normal']; st.font.name = BODY; st.font.size = Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

# footer: sahifa raqami
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, 'ADM MS  ·  Bilim sinovi  ·  ', 8, color=MUTED)
r = fp.add_run(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); r._r.append(fld); r.font.size = Pt(8)

# --- muqova ---
t = table(d, 1, 1, [W]); c = t.cell(0, 0); shade(c, NAVY); borders(c); margins(c, 300, 300, 500, 500)
p = para(c, 6); run(p, 'ADM JIZZAKH  ·  ADM MS', 9, True, GOLD, HEAD)
p = para(c, 0); run(p, 'BILIM', 36, True, 'FFFFFF', HEAD)
p = para(c, 6); run(p, 'SINOVI', 36, True, GOLD, HEAD)
p = para(c, 4); run(p, 'Xodim kompetensiyalarini yozma baholash', 13, False, 'FFFFFF', HEAD)
p = para(c, 0); run(p, 'Evaluation Sheet talablari asosida: xabardorlik, standart ish, koʻrsatkichlar, biznes reja', 9.5, False, 'B9C2D0')
gap(d, 9)

t = table(d, 3, 4, [3.2, 5.3, 3.2, 5.3])
for i, (a, b) in enumerate([('Xodim F.I.Sh', 'Lavozim'), ('Tab №', 'Boʻlim'), ('Sana', 'Baholovchi')]):
    for j, lab in ((0, a), (2, b)):
        lc, vc = t.cell(i, j), t.cell(i, j + 1)
        borders(lc, bottom=(6, LINE)); borders(vc, bottom=(6, LINE)); margins(lc, 90, 40, 0, 60); margins(vc, 90, 40, 60, 60)
        run(lc.paragraphs[0], lab.upper(), 7.5, True, MUTED, HEAD)
gap(d, 9)

t = table(d, 1, 2, [0.25, W - 0.25]); shade(t.cell(0, 0), GOLD); borders(t.cell(0, 0)); c = t.cell(0, 1); shade(c, PALE); borders(c); margins(c, 120, 120, 260, 200)
p = para(c, 4); run(p, 'Qanday ishlash kerak', 11, True, NAVY, HEAD)
for s in ["Har bir savolga oʻz soʻzlaringiz bilan, ish joyingizdan misol keltirib javob yozing.",
          "Hisob-kitobli topshiriqlarda formulani va hisobni toʻliq koʻrsating.",
          "“Vaziyat” topshiriqlarida qadamlaringizni ketma-ket yozing. Tavsiya etilgan vaqt: 60 daqiqa.",
          ]:
    p = para(c, 2); run(p, '▸  ', 10, True, GOLD); run(p, s, 9.5)
gap(d, 9)

p = para(d, 4); run(p, 'NATIJALAR', 9, True, MUTED, HEAD)
rows = [('A', 'Xabardorlik', '11'), ('B', 'Ish joyi va standart ish', '6'), ('C', 'Koʻrsatkichlar', '9'), ('D', 'Biznes reja va takomillashtirish', '7'),
        ('E', 'Tezkor topshiriqlar', '3'), ('', 'Oʻrtacha ball / daraja', '')]
t = table(d, len(rows) + 1, 4, [1.2, 9.3, 3.0, 3.5])
for j, h in enumerate(['', 'Boʻlim', 'Topshiriq', 'Ball (1–5)']):
    c = t.cell(0, j); shade(c, NAVY); borders(c); margins(c, 70, 70); run(c.paragraphs[0], h, 8.5, True, 'FFFFFF', HEAD)
for i, (a, b, n) in enumerate(rows, 1):
    for j, v in enumerate([a, b, n, '']):
        c = t.cell(i, j); borders(c, bottom=(4, LINE)); margins(c, 28, 28)
        if i == len(rows): shade(c, PALE)
        run(c.paragraphs[0], v, 9.5, j == 0 or i == len(rows), GOLD if j == 0 else INK, HEAD if j == 0 else BODY)
gap(d, 6)

legend = [('T', '5', '2E7D5B', 'Mukammal, boshqalarga oʻrgata oladi'), ('D', '4', '4F8A3C', 'Mustaqil, standart va sifat talablariga muvofiq'),
          ('K', '3', 'B88A1E', 'Asosiy bilim bor, yordam yoki nazorat kerak'), ('A', '2', 'C46A2B', 'Umumiy tushuncha, mustaqil qoʻllay olmaydi'),
          ('NT', '1', 'B03A3A', 'Yetarli emas, oʻqitish talab etiladi')]
t = table(d, 1, 5, [W / 5] * 5)
for j, (k, b, col, txt) in enumerate(legend):
    c = t.cell(0, j); borders(c, top=(24, col)); margins(c, 60, 30, 80, 80); t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    p = para(c, 1); run(p, k, 14, True, col, HEAD); run(p, f'  {b} ball', 8, True, MUTED, HEAD)
    p = para(c, 0); run(p, txt, 7.5, color=MUTED)
d.add_page_break()


def banner(letter, title, sub):
    t = table(d, 1, 2, [2.2, W - 2.2])
    a, b = t.cell(0, 0), t.cell(0, 1)
    shade(a, GOLD); shade(b, NAVY); borders(a); borders(b); margins(a, 150, 150, 0, 0); margins(b, 150, 150, 260, 200)
    a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = para(a, 0, WD_ALIGN_PARAGRAPH.CENTER); run(p, letter, 30, True, NAVY, HEAD)
    p = para(b, 0); run(p, title.upper(), 14, True, 'FFFFFF', HEAD)
    p = para(b, 0); run(p, sub, 9, False, 'B9C2D0')
    gap(d, 8)


def answer_box(lines):
    t = table(d, lines, 1, [W])
    for i in range(lines):
        c = t.cell(i, 0); borders(c, bottom=(4, LINE)); margins(c, 0, 0, 120, 120)
        c.paragraphs[0].paragraph_format.line_spacing = Pt(22)
    gap(d, 10)


n = 0
for letter, title, sub, qs in SECTIONS:
    banner(letter, title, sub)
    for kind, tag, q, lines in qs:
        n += 1
        vz = kind == 'v'
        t = table(d, 1, 2, [1.3, W - 1.3]); t.rows[0].cells  # keep together with answer box
        a, b = t.cell(0, 0), t.cell(0, 1)
        shade(a, GOLD if vz else NAVY); borders(a); margins(a, 110, 110, 0, 0); a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(b, 'FBF3E3' if vz else PALE); borders(b); margins(b, 110, 110, 200, 160)
        p = para(a, 0, WD_ALIGN_PARAGRAPH.CENTER); run(p, str(n), 15, True, NAVY if vz else 'FFFFFF', HEAD)
        p = para(b, 2); run(p, tag.upper(), 7.5, True, 'A8741A' if vz else MUTED, HEAD)
        p = para(b, 0); run(p, q, 10.5, True, INK)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
        p = t.cell(0, 1).paragraphs[-1]; p.paragraph_format.keep_with_next = True
        answer_box(lines)
    d.add_page_break()

# ---- E: tezkor topshiriqlar ----
banner('E', 'Tezkor topshiriqlar', 'Toʻgʻri / notoʻgʻri, moslashtirish, hisob jadvali')
p = para(d, 6); run(p, f'{n + 1}.  ', 11, True, GOLD, HEAD); run(p, 'Toʻgʻri yoki notoʻgʻri? Kerakli katakka ✓ belgi qoʻying.', 11, True, NAVY, HEAD)
t = table(d, len(TF) + 1, 4, [0.9, 12.3, 1.9, 1.9])
for j, h in enumerate(['№', 'Fikr', 'Toʻgʻri', 'Notoʻgʻri']):
    c = t.cell(0, j); shade(c, NAVY); borders(c); margins(c, 70, 70); run(c.paragraphs[0], h, 8.5, True, 'FFFFFF', HEAD)
    if j > 1: c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, s in enumerate(TF, 1):
    for j, v in enumerate([str(i), s, '☐', '☐']):
        c = t.cell(i, j); borders(c, bottom=(4, LINE)); margins(c, 90, 90)
        if i % 2 == 0: shade(c, PALE)
        pr = c.paragraphs[0]
        if j > 1: pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(pr, v, 14 if j > 1 else 10, j == 0, GOLD if j == 0 else INK)
gap(d, 16)

p = para(d, 6); run(p, f'{n + 2}.  ', 11, True, GOLD, HEAD); run(p, 'Moslashtiring: har bir qisqartmaga mos taʻrif harfini yozing.', 11, True, NAVY, HEAD)
t = table(d, len(MATCH_L), 4, [2.2, 1.6, 0.8, 12.4])
for i, (l, r_) in enumerate(zip(MATCH_L, MATCH_R)):
    cells = t.rows[i].cells
    shade(cells[0], NAVY); borders(cells[0]); margins(cells[0], 90, 90); run(cells[0].paragraphs[0], l, 10.5, True, 'FFFFFF', HEAD)
    borders(cells[1], bottom=(8, GOLD)); margins(cells[1], 90, 90)
    borders(cells[2]); margins(cells[2], 90, 90, 60, 0); run(cells[2].paragraphs[0], 'abcdef'[i] + ')', 10, True, GOLD, HEAD)
    borders(cells[3], bottom=(4, LINE)); margins(cells[3], 90, 90); run(cells[3].paragraphs[0], r_, 10)
gap(d, 16)

d.add_page_break(); p = para(d, 6); run(p, f'{n + 3}.  ', 11, True, GOLD, HEAD); run(p, 'Hisob jadvali: formulani yozing va natijani toping.', 11, True, NAVY, HEAD)
t = table(d, len(CALC) + 1, 4, [3.4, 4.6, 5.4, 3.6])
for j, h in enumerate(['Koʻrsatkich', 'Berilgan', 'Formula va hisob', 'Javob']):
    c = t.cell(0, j); shade(c, NAVY); borders(c); margins(c, 70, 70); run(c.paragraphs[0], h, 8.5, True, 'FFFFFF', HEAD)
for i, (k, g) in enumerate(CALC, 1):
    for j, v in enumerate([k, g, '', '']):
        c = t.cell(i, j); borders(c, bottom=(4, LINE), right=(4, LINE) if j < 3 else None); margins(c, 200, 200)
        if j == 3: shade(c, 'FBF3E3')
        run(c.paragraphs[0], v, 10, j == 0, NAVY if j == 0 else INK, HEAD if j == 0 else BODY)
gap(d, 18)

t = table(d, 1, 1, [W]); c = t.cell(0, 0); shade(c, NAVY); borders(c); margins(c, 220, 220, 300, 300)
p = para(c, 2, WD_ALIGN_PARAGRAPH.CENTER); run(p, 'Javoblaringizni Claudeʻga yuboring', 12, True, GOLD, HEAD)
p = para(c, 0, WD_ALIGN_PARAGRAPH.CENTER); run(p, 'Har bir javob T / D / K / A / NT shkalasi boʻyicha baholanadi va nima yetishmayotgani tushuntiriladi.', 9.5, False, 'FFFFFF')

d.save('Bilim_sinovi.docx')
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', 'Bilim_sinovi.docx'], check=True, capture_output=True)
print('savollar:', n + 3)
