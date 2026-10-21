# Evaluation sheet (Abdumalikov Ulug'bek) asosida yozma savolnoma: savol + javob uchun bo'sh qatorlar.
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

S = [
 ("A. Talab etiladigan xabardorlik darajasi", [
  ("PE-1", "Korxonaning qadriyatlari, missiyasi va strategik maqsadlari kundalik ishingizda qanday namoyon bo'ladi? Ikkita amaliy misol keltiring.", 5),
  ("PE-7", "R&R (Rol va Javobgarliklar) nima va ular nima uchun tashkilotning har bir darajasi uchun belgilanadi? R&R ADM MS ning nechta tamoyilini qamrab olishi kerak?", 5),
  ("PE-7", "O'z lavozimingiz bo'yicha asosiy majburiyatlaringiz, huquqlaringiz va javobgarliklaringizni yozing.", 7),
  ("PE-7", "Guruh rahbari (GL) uchun belgilangan 4 ta asosiy vazifani sanab bering.", 5),
  ("PE-2", "Ish joyi xavfsizligi tizimi (WSS) nimaga xizmat qiladi? \"Har bir Inson, Har bir Ob'ekt, Har Kuni\" tamoyilini qanday tushunasiz?", 5),
  ("PE-2", "Zavodda risklarni baholash uchun qaysi vosita qo'llanilishi shart? ESCP eskalatsiya tartibida muammo 24 soatda, 7 kunda va 30 kunda hal etilmasa nima qilinadi?", 6),
  ("PE-2", "Xavfsizlik aylanmalari (safety walk) qanday turlarga bo'linadi? Sexda rioya qilinishi shart bo'lgan tanqidiy xavfsizlik tashabbuslariga 3 ta misol keltiring.", 6),
  ("PE-3", "GCTA (Global Call to Actions) nima? SIF qanday hodisa? Ular bo'yicha tuzatuvchi chora-tadbirlar rejasi (CAP) qanday elementlarni o'z ichiga olishi kerak va qanday tasdiqlanadi?", 7),
  ("PE-6", "JIT (Job Instruction Training) usuli nima va kim tomonidan qo'llaniladi? Skill flexibility/versatility chart (4 kvadrant) nimani ko'rsatadi va xodim operatsiyani 3 oy bajarmasa nima o'zgaradi?", 6),
  ("Intizom", "Korxonaning ichki tartib-qoidalari va mehnat intizomi bo'yicha xodimdan talab qilinadigan asosiy qoidalarni yozing.", 6),
 ]),
 ("B. Ish joyini tashkil etish va standartlashtirilgan ish", [
  ("ST-3", "5S bosqichlarini tartib bilan yozing va har birini bir jumlada izohlang.", 6),
  ("ST-3", "\"Red tag\" jarayoni va \"shadow board\" konsepsiyasi nima uchun kerak? Doskalarda holat qanday umumiy belgi bilan ko'rsatiladi?", 5),
  ("ST-3", "7 turdagi yo'qotishni sanab bering. WPO orqali ulardan birini bartaraf etishga misol keltiring.", 6),
  ("ST-6", "SOS, JES, STS va TIS hujjatlari nima? Qaysilari siklik, qaysilari siklik bo'lmagan ish uchun?", 6),
  ("ST-6", "Siklik standartlashtirilgan ish hujjatida qanday ma'lumotlar bo'lishi kerak (kamida 5 ta)? Ish elementlari sikl vaqtining taxminan necha foizi hajmida bo'laklanadi?", 6),
  ("ST-8", "Operator standartlashtirilgan ishga amal qilayotganini qanday tekshirasiz? SOSdagi umumiy sikl vaqti real holatdan necha foizdan ko'p farq qilmasligi kerak?", 6),
 ]),
 ("C. Ishlab chiqarish ko'rsatkichlari (hisob-kitob bilan)", [
  ("ST-4", "Takt Time nima va qanday hisoblanadi? Masala: smena 8 soat, tanaffuslar jami 30 daqiqa, smenalik reja 225 avtomobil. Takt Time necha soniya?", 6),
  ("Tahlil", "Ishlab chiqarish rejasidagi avtomobillarni Takt Time bo'yicha tahlil qilishning 4 ta afzalligini yozing.", 5),
  ("Tahlil", "Cycle Time va Line Balance nima? Masala: Cycle Time = 120 soniya, bitta stansiyaning operatsiyasi 150 soniya. Qanday muammo yuzaga keladi va uning oqibatlari qanday?", 6),
  ("Tahlil", "HPV qanday hisoblanadi? Masala: oyda 4 800 soat ishlangan, 300 avtomobil chiqarilgan. HPV nechaga teng? HPV 18 soatdan 16 soatga tushsa, samaradorlik necha foiz yaxshilangan?", 6),
  ("Tahlil", "HPV oshishiga qaysi omillar sabab bo'ladi (kamida 5 ta)? Ishlab chiqarish hajmi oshganda HPV odatda qanday o'zgaradi va nega?", 6),
  ("Tahlil", "UPH qanday hisoblanadi? Masala: liniya 8 soatda 450 avtomobil chiqardi. UPH nechaga teng? Target 60, Actual 55 UPH bo'lsa, bu nimani bildiradi?", 5),
  ("Downtime", "Rejalashtirilgan va rejalashtirilmagan Downtime farqi nima? Har biriga 3 tadan misol keltiring. Downtime nima uchun o'lchanadi?", 6),
  ("Tahlil", "I/Ch sexlarining absentizm ko'rsatkichi nima uchun jamlanib rahbariyatga taqdim etiladi? Kamida 4 ta sabab yozing.", 5),
  ("Tahlil", "#630, #690, #700 eventlar tahlili nimaga xizmat qiladi? \"04\" avtomobillar nima va ular bo'yicha qanday choralar ko'rasiz?", 6),
  ("Tahlil", "Oylik ishlab chiqarilgan avtomobillarni brend va model kesimida tahlil qilish qaysi 4 yo'nalishda foyda beradi?", 6),
 ]),
 ("D. Biznes rejani boshqarish va takomillashtirish", [
  ("SPQRC", "SPQRC(E) yo'nalishlarining har bir harfi nimani bildiradi va ular kundalik boshqaruvda nima uchun muhim?", 5),
  ("CI-1", "BPD nima? Yagona BPD formati qanday elementlarni o'z ichiga oladi va korxonada odatda nechta BPD darajasi bo'ladi?", 6),
  ("CI-2", "SMART tamoyilining har bir harfini izohlang. Metrika doimiy \"yashil\" yoki doimiy \"qizil\" bo'lsa, nima qilish kerak?", 6),
  ("PDCA", "PDCA siklining bosqichlarini yozing va o'z ishingizdagi bitta muammoni shu sikl bo'yicha qanday hal qilishingizni ko'rsating.", 7),
  ("Q-10 / PE-15/16", "Error Proofing (xatolarning oldini olish) nima va uning validatsiyasi qanday o'tkaziladi? \"Go and See\" va Layered Audit nima?", 6),
  ("MS Office", "Excel'da reja-fakt tahlili uchun qaysi formulalardan foydalanasiz? Ikkita misol yozing.", 4),
  ("Baholash", "Baholash shkalasidagi T, D, K, A va NT darajalari nimani bildiradi? Yakuniy natija (Result) qanday hisoblanadi?", 5),
 ]),
]

d = Document()
sec = d.sections[0]; sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run("YOZMA SAVOLNOMA"); r.bold = True; r.font.size = Pt(16)
h2 = d.add_paragraph(); h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h2.add_run("Evaluation Sheet talablari asosida · ADM MS"); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
for lab in ("Xodim F.I.Sh: ______________________________", "Lavozim: ______________________   Sana: ____________"):
    d.add_paragraph(lab)
p = d.add_paragraph(); r = p.add_run("Ko'rsatma: har bir savolga o'z so'zlaringiz bilan javob yozing. Hisob-kitobli savollarda formulani va hisobni ko'rsating.")
r.italic = True; r.font.size = Pt(10)
n = 0
for title, qs in S:
    t = d.add_paragraph(); t.paragraph_format.space_before = Pt(14)
    r = t.add_run(title); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
    for tag, q, lines in qs:
        n += 1
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{n}. "); r.bold = True
        r = p.add_run(f"[{tag}] "); r.font.color.rgb = RGBColor(0x88, 0x66, 0x22); r.font.size = Pt(9)
        p.add_run(q).bold = True
        a = d.add_paragraph(); a.add_run("Javob:").italic = True
        for _ in range(lines):
            l = d.add_paragraph("_" * 86); l.paragraph_format.space_after = Pt(2)
            l.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
d.save('Yozma_savolnoma.docx')
# chatda javob yozish uchun matn varianti
with open('Yozma_savolnoma.txt', 'w', encoding='utf-8') as f:
    n = 0
    for title, qs in S:
        f.write(f"\n{title}\n")
        for tag, q, _ in qs:
            n += 1; f.write(f"\n{n}. [{tag}] {q}\nJavob: \n")
print(n, 'savol')
